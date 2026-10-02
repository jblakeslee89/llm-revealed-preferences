"""Build briefings/project-overview-5min.pptx from the 12-minute Phase 3 deck.

Keeps slides 1, 3, 5, 7, 9, 10 and 11 of phase3-findings-12min.pptx (same RAND purple
layouts and figures), rewrites their text and speaker notes as a whole-project overview,
updates the footer band and cuts the progress dots to seven.

Usage:
    python briefings/build_overview_5min.py
"""

from __future__ import annotations

import copy
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent
SRC = HERE / "phase3-findings-12min.pptx"
OUT = HERE / "project-overview-5min.pptx"
KEEP = [1, 3, 5, 7, 9, 10, 11]  # 1-based slide numbers in the source deck
EMU = 914400

FOOTER = {"Phase 3 findings": "Project overview", "September 2026": "October 2026",
          "12 slides": "7 slides"}

# per output slide: {shape name: new text}, then speaker notes
TEXT = [
    {
        "Text 2": "We put that question to nine AI models from four families, 42,505 times in all.",
        "Text 3": "Their choices have a structure, it is mostly set at one training step, and it changes when they think first.",
    },
    {
        "Text 24": "We read each answer straight out of the model's internal probabilities. Since September, "
                   "the models may also reason in writing before they answer.",
        "Text 30": "9",
        "Text 31": "models",
        "Text 36": "42,505",
    },
    {},
    {},
    {},
    {
        "Text 1": "The written verdict can contradict the arithmetic",
        "Text 22": "Right sums, wrong verdict",
        "Text 23": "Most of Qwen’s remaining errors compute the expected value correctly and then write "
                   "“$152.04, which is less than the guaranteed $117.”",
        "Text 25": "Qwen leans safe",
        "Text 26": "Its backward verdicts all favor the sure amount and follow one stock phrase. Handed the "
                   "right number, after “which is” it says “less” 89% of the time.",
        "Text 28": "Llama leans the other way",
        "Text 29": "Llama writes “Since $115.72 is greater than $149,” mostly when the gamble is listed "
                   "first. Each family has its own habit.",
        "Text 31": "Telling works, paying did not",
        "Text 32": "Told the rule, Qwen stops writing backward verdicts. Paid through training for snap "
                   "answers, a model learned one fixed answer, never the arithmetic.",
    },
    {
        "Text 22": "Look at fine-tuning",
        "Text 23": "Most of the snap-answer profile is in place after supervised fine-tuning, the first "
                   "assistant-training step. An audit that starts at preference training starts late.",
        "Text 25": "Say which mode you measured",
        "Text 26": "Preference training reshapes how a model reasons and leaves its snap-answer preferences where "
                   "they were. A bias found in one mode can vanish or reverse in the other.",
        "Text 28": "Check that the subject is awake",
        "Text 29": "A model can return confident answers while failing to prefer $70 to $50. Eighty free-money "
                   "questions catch it, and cost almost nothing to run.",
        "Text 31": "What comes next",
        "Text 32": "Pay the model while it reasons. Ask whether models favor their own family, and whether "
                   "these risk habits carry into the advice they give commanders and officials.",
    },
]

NOTES = [
    "A sure eighty-six dollars, or a forty-three percent shot at two hundred and ten? Economists use "
    "choices like this to measure how people handle risk. Since July I have asked AI models the same "
    "kind of question, more than forty-two thousand times, across nine models from four families. "
    "Three things came out of it. The models' choices have a measurable structure. Most of that "
    "structure is set at one early training step. And it changes when the model is allowed to think "
    "before it answers.",

    "The method is simple. Every model sees the same sixteen hundred and eighty questions, and only one "
    "thing changes at a time: the odds, the prize, the wording, or the order of the options. We read "
    "the answer from the model's internal probabilities, so each question is asked once and the answer "
    "is exact. Eighty questions are free money, a sure fifty against a sure seventy, to check the model "
    "is paying attention. Since September we also let the models write out their reasoning first.",

    "OLMo publishes a snapshot after every training step, so we could ask the same questions at four "
    "points in one model's life. Most of the change happens at supervised fine-tuning, the first and "
    "simplest step, where the model learns from examples of helpful answers. Free-money accuracy jumps "
    "from fifty-six to seventy-six percent there, the largest single step, and the preferences we "
    "estimate stop moving after it. Preference training, the step usually blamed for human-like biases, "
    "finds the profile already in place. When a model answers at once, its money personality is mostly "
    "set early.",

    "The second finding is that the bias belongs to the training recipe. Describe the same bet as a "
    "loss instead of a smaller gain, with identical arithmetic. Claude Haiku gambles less, Qwen gambles "
    "far more, OLMo somewhat more, and Llama barely reacts. Four systems, four reactions. A framing "
    "weakness documented on one model says little about another.",

    "Third, all of that describes snap answers. Told to maximize expected value and answering at once, "
    "the models we tested sit near a coin flip. Allowed to reason first, they follow the rule seventy-one to "
    "ninety-six percent of the time, and the loss-wording effect shrinks, vanishes, or reverses. In "
    "OLMo this change happens at preference training: the step that left the snap answers alone is the "
    "one that changes how the model reasons.",

    "Reading the reasoning brought a surprise. Most remaining errors get the arithmetic right and then "
    "state the comparison backward. Qwen's backward verdicts always favor the sure amount and follow one "
    "phrase, which is less than the guaranteed amount. Llama's go the other way and favor whichever "
    "option it read first. Telling the model the rule removes Qwen's errors. Paying a model through "
    "training for its snap answers taught it a fixed answer and nothing else.",

    "Four takeaways. Audit at the fine-tuning checkpoint, because that is where most of the snap-answer "
    "profile appears. Say whether you measured snap answers or reasoning, because the two can disagree. "
    "Check that the subject is paying attention before trusting any estimate. And read reasoning traces "
    "knowing each family has its own habits. Next, we pay the model while it reasons, use the same "
    "machinery to ask whether AI agents favor their own kind, and then ask whether the risk habits we "
    "measured show up in the advice a model gives a commander or an official, with the same numbers as "
    "our bets. Thank you.",
]


def set_text(tf, new):
    """Replace a text frame's text, keeping the first run's formatting."""
    paras = tf.paragraphs
    runs = paras[0].runs
    runs[0].text = new
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for p in paras[1:]:
        p._p.getparent().remove(p._p)


def drop_slides(prs, keep):
    ids = prs.slides._sldIdLst
    for i, sld in reversed(list(enumerate(list(ids), start=1))):
        if i not in keep:
            prs.part.drop_rel(sld.rId)
            ids.remove(sld)


def fix_pips(slide, active, n=7):
    pips = sorted([s for s in slide.shapes if s.top is not None and abs(s.top / EMU - 10.42) < 0.05
                   and abs(s.width / EMU - 0.2) < 0.02], key=lambda s: s.left)
    for s in pips[:-n]:
        s._element.getparent().remove(s._element)
    pips = pips[-n:]
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for k, s in enumerate(pips):
        clr = s._element.spPr.find("a:solidFill/a:srgbClr", ns)
        for a in clr.findall("a:alpha", ns):
            clr.remove(a)
        if k != active:
            alpha = copy.deepcopy(clr.makeelement("{%s}alpha" % ns["a"], {"val": "22000"}))
            clr.append(alpha)


def main():
    prs = Presentation(SRC)
    drop_slides(prs, KEEP)
    assert len(prs.slides) == 7
    for k, slide in enumerate(prs.slides):
        by_name = {s.name: s for s in slide.shapes if s.has_text_frame}
        for s in by_name.values():
            t = s.text_frame.text
            if t in FOOTER:
                set_text(s.text_frame, FOOTER[t])
        for name, new in TEXT[k].items():
            set_text(by_name[name].text_frame, new)
        set_text(slide.notes_slide.notes_text_frame, NOTES[k])
        fix_pips(slide, k)
    prs.save(OUT)
    words = sum(len(n.split()) for n in NOTES)
    print(f"wrote {OUT.name}: {len(prs.slides)} slides, {words} words of notes (~{words / 140:.1f} min)")


if __name__ == "__main__":
    main()
